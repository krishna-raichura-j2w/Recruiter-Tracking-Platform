"""
Tickets Excel export — three-sheet workbook:
  Sheet 1: Ticket Summary      (one row per ticket, all key fields)
  Sheet 2: Hierarchy Details   (one row per step per ticket)
  Sheet 3: Comments            (one row per comment per ticket)
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from features.hrbp.tickets.service import _enrich_ticket
from features.hrbp.storage.service import upload_bytes
from infra.hrbp_models import HRBPTicket
from infra.models import User
from sqlalchemy import Text, cast, or_

# ── Style constants ────────────────────────────────────────────────────────────

_THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

BODY_FONT   = Font(name="Calibri", size=10)
BODY_BOLD   = Font(name="Calibri", size=10, bold=True)
ALIGN_CTR   = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT  = Alignment(horizontal="left",   vertical="center", wrap_text=True)
ALIGN_WRAP  = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

FILLS = {
    "header_blue":   PatternFill(fill_type="solid", fgColor="1A56DB"),
    "header_teal":   PatternFill(fill_type="solid", fgColor="0D7490"),
    "header_purple": PatternFill(fill_type="solid", fgColor="5521B5"),
    "open":          PatternFill(fill_type="solid", fgColor="D1FAE5"),
    "closed":        PatternFill(fill_type="solid", fgColor="F3F4F6"),
    "escalated":     PatternFill(fill_type="solid", fgColor="FEE2E2"),
    "critical":      PatternFill(fill_type="solid", fgColor="FCA5A5"),
    "high":          PatternFill(fill_type="solid", fgColor="FED7AA"),
    "medium":        PatternFill(fill_type="solid", fgColor="FEF9C3"),
    "low":           PatternFill(fill_type="solid", fgColor="DBEAFE"),
    "breached":      PatternFill(fill_type="solid", fgColor="FEE2E2"),
    "on_track":      PatternFill(fill_type="solid", fgColor="D1FAE5"),
}

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)


def _hdr(ws, col: int, row: int, value: str, fill: PatternFill, width: float):
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill = fill
    cell.font = HEADER_FONT
    cell.alignment = ALIGN_CTR
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(col)].width = width


def _cell(ws, row: int, col: int, value, align=ALIGN_LEFT, bold=False, fill=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = BODY_BOLD if bold else BODY_FONT
    cell.alignment = align
    cell.border = BORDER
    if fill:
        cell.fill = fill
    return cell


def _strip_html(html: str | None) -> str:
    """Remove HTML tags and decode basic entities for plain-text cells."""
    if not html:
        return ""
    text = re.sub(r"<[^>]+>", " ", html)
    text = text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&nbsp;", " ")
    return re.sub(r" {2,}", " ", text).strip()


def _fmt(dt) -> str:
    if not dt:
        return ""
    if isinstance(dt, str):
        try:
            dt = datetime.fromisoformat(dt.replace("Z", "+00:00"))
        except Exception:
            return dt
    return dt.strftime("%-d %b %Y %I:%M %p")


def _sla_status(ticket_dict: dict) -> str:
    status = ticket_dict.get("status", "")
    if status == "closed":
        return "Closed"
    deadline = ticket_dict.get("sla_deadline")
    if not deadline:
        return "No SLA"
    try:
        if isinstance(deadline, str):
            deadline = datetime.fromisoformat(deadline.replace("Z", "+00:00"))
        now = datetime.now(timezone.utc)
        deadline = deadline if deadline.tzinfo else deadline.replace(tzinfo=timezone.utc)
        return "Breached" if now > deadline else "On Track"
    except Exception:
        return "—"


# ── Sheet 1: Ticket Summary ────────────────────────────────────────────────────

_S1_HEADERS = [
    ("TICKET #",             12),
    ("TITLE",                32),
    ("SOP TYPE",             12),
    ("SOP NAME",             28),
    ("STATUS",               12),
    ("PRIORITY",             12),
    ("SLA STATUS",           13),
    ("SLA DEADLINE",         20),
    ("PO AT RISK (₹)",       16),
    ("CLIENT",               22),
    ("RAISED BY",            20),
    ("ESCALATION MGR",       20),
    ("CONSULTANTS",          32),
    ("CONSULTANT COUNT",     16),
    ("CURRENT STEP #",       14),
    ("CURRENT STEP LABEL",   24),
    ("CURRENT STEP OWNER",   22),
    ("TOTAL STEPS",          12),
    ("STEP SLA EXTENDED",    22),
    ("ATTACHMENTS",          12),
    ("COMMENTS",             10),
    ("DESCRIPTION",          48),
    ("CREATED AT",           20),
    ("CLOSED AT",            20),
    ("LAST UPDATED",         20),
]


def _build_summary(ws, tickets: list[dict]):
    ws.title = "Ticket Summary"
    ws.row_dimensions[1].height = 24
    fill = FILLS["header_blue"]
    for col_idx, (hdr, width) in enumerate(_S1_HEADERS, start=1):
        _hdr(ws, col_idx, 1, hdr, fill, width)

    for row_idx, t in enumerate(tickets, start=2):
        ws.row_dimensions[row_idx].height = 18

        hierarchy = t.get("hierarchy_json") or []
        step_idx  = (t.get("current_step") or 1) - 1
        cur_step  = hierarchy[step_idx] if 0 <= step_idx < len(hierarchy) else {}

        consultants = t.get("consultants") or []
        consultant_names = ", ".join(c["name"] for c in consultants)

        sla_stat = _sla_status(t)
        status   = (t.get("status") or "").lower()
        priority = (t.get("priority") or "").lower()

        po = t.get("po_risk_amount")
        po_str = f"₹{po:,.0f}" if po else "—"

        step_ext = _fmt(t.get("step_sla_extended_until"))

        values = [
            t.get("ticket_number", ""),
            t.get("title", ""),
            t.get("sop_type") or "—",
            t.get("sop_name") or "—",
            status.capitalize(),
            priority.capitalize(),
            sla_stat,
            _fmt(t.get("sla_deadline")),
            po_str,
            t.get("client_name") or "—",
            t.get("raised_by_name") or "—",
            t.get("escalation_mgr_name") or "—",
            consultant_names or "—",
            len(consultants),
            t.get("current_step", 1),
            cur_step.get("label", "—"),
            cur_step.get("user_name") or "—",
            len(hierarchy),
            step_ext or "—",
            len(t.get("attachments") or []),
            len(t.get("comments") or []),
            _strip_html(t.get("description")),
            _fmt(t.get("created_at")),
            _fmt(t.get("closed_at")),
            _fmt(t.get("updated_at")),
        ]

        for col_idx, value in enumerate(values, start=1):
            align = ALIGN_WRAP if col_idx == 22 else (ALIGN_CTR if col_idx in (5, 6, 7, 14, 15, 17, 18, 20, 21) else ALIGN_LEFT)
            cell_fill = None

            if col_idx == 5:   # Status
                cell_fill = FILLS.get(status)
            elif col_idx == 6:  # Priority
                cell_fill = FILLS.get(priority)
            elif col_idx == 7:  # SLA Status
                if sla_stat == "Breached":
                    cell_fill = FILLS["breached"]
                elif sla_stat == "On Track":
                    cell_fill = FILLS["on_track"]

            c = _cell(ws, row_idx, col_idx, value, align=align, fill=cell_fill,
                      bold=(col_idx == 1))

            if col_idx == 9 and po:  # PO at Risk
                c.font = Font(name="Calibri", size=10, bold=True, color="DC2626")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_S1_HEADERS))}1"


# ── Sheet 2: Hierarchy Details ─────────────────────────────────────────────────

_S2_HEADERS = [
    ("TICKET #",       12),
    ("TITLE",          28),
    ("STATUS",         12),
    ("STEP ORDER",     11),
    ("STEP LABEL",     24),
    ("ROLE",           14),
    ("ASSIGNEE",       22),
    ("SLA WINDOW",     18),
    ("SLA HOURS",      11),
    ("RESOLVED AT",    22),
    ("RESOLVED BY",    22),
    ("IS CURRENT",     12),
]


def _build_hierarchy(ws, tickets: list[dict]):
    ws.title = "Hierarchy Details"
    ws.row_dimensions[1].height = 24
    fill = FILLS["header_teal"]
    for col_idx, (hdr, width) in enumerate(_S2_HEADERS, start=1):
        _hdr(ws, col_idx, 1, hdr, fill, width)

    row_idx = 2
    for t in tickets:
        hierarchy = t.get("hierarchy_json") or []
        current_step = t.get("current_step", 1)
        for step in hierarchy:
            is_current = (step.get("order") == current_step) and t.get("status") == "open"
            values = [
                t.get("ticket_number", ""),
                t.get("title", ""),
                (t.get("status") or "").capitalize(),
                step.get("order", ""),
                step.get("label", ""),
                step.get("role", ""),
                step.get("user_name") or "—",
                step.get("sla_window") or "—",
                step.get("sla_hours") or "—",
                _fmt(step.get("resolved_at")),
                step.get("resolved_by_name") or "—",
                "Yes" if is_current else "",
            ]
            ws.row_dimensions[row_idx].height = 16
            for col_idx, value in enumerate(values, start=1):
                align = ALIGN_CTR if col_idx in (3, 4, 8, 9, 12) else ALIGN_LEFT
                cell_fill = FILLS["header_teal"] if is_current and col_idx == 12 else None
                _cell(ws, row_idx, col_idx, value, align=align, fill=cell_fill,
                      bold=(col_idx == 1 or is_current))
            row_idx += 1

    ws.freeze_panes = "A2"


# ── Sheet 3: Comments ──────────────────────────────────────────────────────────

_S3_HEADERS = [
    ("TICKET #",        12),
    ("TITLE",           28),
    ("AUTHOR",          22),
    ("STEP #",          8),
    ("IS RESOLUTION",   14),
    ("COMMENT",         60),
    ("POSTED AT",       20),
]


def _build_comments(ws, tickets: list[dict]):
    ws.title = "Comments"
    ws.row_dimensions[1].height = 24
    fill = FILLS["header_purple"]
    for col_idx, (hdr, width) in enumerate(_S3_HEADERS, start=1):
        _hdr(ws, col_idx, 1, hdr, fill, width)

    row_idx = 2
    for t in tickets:
        for cm in (t.get("comments") or []):
            is_res = cm.get("is_resolution", False)
            values = [
                t.get("ticket_number", ""),
                t.get("title", ""),
                cm.get("author_name") or "—",
                cm.get("hierarchy_step", ""),
                "Yes" if is_res else "No",
                _strip_html(cm.get("content", "")),
                _fmt(cm.get("created_at")),
            ]
            ws.row_dimensions[row_idx].height = 30
            for col_idx, value in enumerate(values, start=1):
                align = ALIGN_WRAP if col_idx == 6 else (ALIGN_CTR if col_idx in (4, 5) else ALIGN_LEFT)
                res_fill = PatternFill(fill_type="solid", fgColor="D1FAE5") if is_res else None
                _cell(ws, row_idx, col_idx, value, align=align,
                      fill=res_fill, bold=(col_idx == 1))
            row_idx += 1

    ws.freeze_panes = "A2"


# ── Public entry point ─────────────────────────────────────────────────────────

def build_and_upload(
    db: Session,
    current_user: User,
    status: str | None = None,
    priority: str | None = None,
    client_id: int | None = None,
    sop_id: int | None = None,
    search: str | None = None,
) -> str:
    q = db.query(HRBPTicket)

    role = current_user.role.value if hasattr(current_user.role, "value") else str(current_user.role)
    if role in ("hrbp", "bh"):
        q = q.filter(
            or_(
                HRBPTicket.raised_by_id == current_user.id,
                HRBPTicket.escalation_mgr_id == current_user.id,
                cast(HRBPTicket.hierarchy_json, Text).contains(str(current_user.id)),
            )
        )

    if status:
        q = q.filter(HRBPTicket.status == status)
    if priority:
        q = q.filter(HRBPTicket.priority == priority)
    if client_id:
        q = q.filter(HRBPTicket.client_id == client_id)
    if sop_id:
        q = q.filter(HRBPTicket.sop_id == sop_id)
    if search:
        q = q.filter(HRBPTicket.ticket_number.ilike(f"%{search}%"))

    q = q.order_by(HRBPTicket.created_at.desc())
    raw_tickets = q.all()

    tickets = [_enrich_ticket(db, t) for t in raw_tickets]

    wb = Workbook()
    ws1 = wb.active
    ws2 = wb.create_sheet()
    ws3 = wb.create_sheet()

    _build_summary(ws1, tickets)
    _build_hierarchy(ws2, tickets)
    _build_comments(ws3, tickets)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return upload_bytes(buf.read(), f"tickets_data_{timestamp}.xlsx")
