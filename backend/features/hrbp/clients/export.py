"""
Clients Excel export — single-sheet workbook with all client data.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session, aliased

from features.hrbp.storage.service import upload_bytes
from infra.hrbp_models import HRBPClient, HRBPConsultant
from infra.models import User

_THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
BODY_FONT = Font(name="Calibri", size=10)
BODY_BOLD = Font(name="Calibri", size=10, bold=True)
ALIGN_CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1A56DB")
ACTIVE_FILL = PatternFill(fill_type="solid", fgColor="D1FAE5")
INACTIVE_FILL = PatternFill(fill_type="solid", fgColor="F3F4F6")


def _hdr(ws, col: int, value: str, width: float):
    cell = ws.cell(row=1, column=col, value=value)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = ALIGN_CTR
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(col)].width = width


def _cell(ws, row: int, col: int, value, align=ALIGN_LEFT, bold=False, fill=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = BODY_BOLD if bold else BODY_FONT
    cell.alignment = align
    cell.border = BORDER
    if fill:
        cell.fill = fill
    return cell


def _fmt(dt) -> str:
    if not dt:
        return ""
    if isinstance(dt, str):
        try:
            dt = datetime.fromisoformat(dt.replace("Z", "+00:00"))
        except Exception:
            return dt
    return dt.strftime("%-d %b %Y")


_HEADERS = [
    ("CLIENT NAME",       28),
    ("INDUSTRY",          18),
    ("HRBP",              30),
    ("BH OWNER",          22),
    ("HEADCOUNT",         12),
    ("MONTHLY PO (₹)",    18),
    ("STATUS",            12),
    ("CREATED AT",        16),
]


def build_and_upload(
    db: Session,
    current_user: User,
    search: str | None = None,
    industry: str | None = None,
    is_active: bool | None = None,
) -> str:
    role = current_user.role.value if hasattr(current_user.role, "value") else str(current_user.role)

    # Build id → name lookup for all HRBP users so we can resolve hrbp_ids array
    hrbp_name_map: dict[int, str] = {
        u.id: u.name
        for u in db.query(User.id, User.name).filter(User.role == "hrbp").all()
    }

    consultant_sub = (
        db.query(
            HRBPConsultant.client_id,
            func.count(HRBPConsultant.id).label("headcount"),
            func.coalesce(func.sum(HRBPConsultant.monthly_po), 0).label("total_monthly_po"),
        )
        .group_by(HRBPConsultant.client_id)
        .subquery()
    )

    BhUser = aliased(User)

    q = (
        db.query(
            HRBPClient.name,
            HRBPClient.industry,
            HRBPClient.hrbp_id,
            HRBPClient.hrbp_ids,
            BhUser.name.label("bh_name"),
            func.coalesce(consultant_sub.c.headcount, 0).label("headcount"),
            func.coalesce(consultant_sub.c.total_monthly_po, 0).label("total_monthly_po"),
            HRBPClient.is_active,
            HRBPClient.created_at,
        )
        .outerjoin(BhUser, HRBPClient.bh_id == BhUser.id)
        .outerjoin(consultant_sub, HRBPClient.id == consultant_sub.c.client_id)
    )

    if role == "hrbp":
        uid = current_user.id
        q = q.filter(
            or_(
                HRBPClient.hrbp_ids.contains([uid]),
                and_(
                    func.coalesce(func.array_length(HRBPClient.hrbp_ids, 1), 0) == 0,
                    HRBPClient.hrbp_id == uid,
                ),
            )
        )
    elif role == "bh":
        q = q.filter(HRBPClient.bh_id == current_user.id)

    if search:
        q = q.filter(HRBPClient.name.ilike(f"%{search}%"))
    if industry:
        q = q.filter(HRBPClient.industry.ilike(f"%{industry}%"))
    if is_active is not None:
        q = q.filter(HRBPClient.is_active == is_active)

    q = q.order_by(HRBPClient.name)
    rows = q.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Clients"
    ws.row_dimensions[1].height = 24

    for col_idx, (hdr, width) in enumerate(_HEADERS, start=1):
        _hdr(ws, col_idx, hdr, width)

    for row_idx, r in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 18
        status = "Active" if r.is_active else "Inactive"
        status_fill = ACTIVE_FILL if r.is_active else INACTIVE_FILL
        po = float(r.total_monthly_po or 0)

        # Resolve HRBP names: use hrbp_ids array if populated, else fall back to hrbp_id
        ids = r.hrbp_ids if r.hrbp_ids else ([r.hrbp_id] if r.hrbp_id else [])
        hrbp_display = ", ".join(hrbp_name_map.get(i, f"#{i}") for i in ids) or "—"

        values = [
            r.name or "—",
            r.industry or "—",
            hrbp_display,
            r.bh_name or "—",
            r.headcount or 0,
            f"₹{po:,.0f}" if po else "—",
            status,
            _fmt(r.created_at),
        ]
        for col_idx, value in enumerate(values, start=1):
            fill = status_fill if col_idx == 7 else None
            align = ALIGN_CTR if col_idx in (5, 7, 8) else ALIGN_LEFT
            _cell(ws, row_idx, col_idx, value, align=align, bold=(col_idx == 1), fill=fill)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}1"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return upload_bytes(buf.read(), "clients_report.xlsx")
