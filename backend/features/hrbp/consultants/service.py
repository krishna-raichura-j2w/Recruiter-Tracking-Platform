from core.pagination import PageResult, paginate
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from fastapi import HTTPException
from infra.hrbp_models import HRBPClient, HRBPConsultant, HRBPTicket, hrbp_ticket_consultants
from infra.models import User
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from features.hrbp.consultants.schema import ConsultantCreate, ConsultantUpdate


def get_summary(db: Session, current_user: User, client_id: int | None = None) -> dict:
    role = current_user.role.value if hasattr(current_user.role, "value") else str(current_user.role)

    q = db.query(HRBPConsultant)
    if role == "hrbp":
        q = q.filter(HRBPConsultant.hrbp_id == current_user.id)
    elif role == "bh":
        bh_client_ids = [
            r.id for r in db.query(HRBPClient.id).filter_by(bh_id=current_user.id).all()
        ]
        q = q.filter(HRBPConsultant.client_id.in_(bh_client_ids))

    if client_id is not None:
        q = q.filter(HRBPConsultant.client_id == client_id)

    total  = q.count()
    active = q.filter(HRBPConsultant.is_active.is_(True)).count()

    cutoff = date.today() + timedelta(days=30)
    expiring_soon = q.filter(
        HRBPConsultant.is_active.is_(True),
        HRBPConsultant.po_end_date <= cutoff,
        HRBPConsultant.po_end_date >= date.today(),
    ).count()

    po_at_risk = q.filter(
        HRBPConsultant.is_active.is_(True),
        HRBPConsultant.po_risk > 0,
    ).count()

    clients_served = (
        q.with_entities(func.count(func.distinct(HRBPConsultant.client_id)))
        .filter(HRBPConsultant.client_id.isnot(None))
        .scalar()
        or 0
    )

    return {
        "total":          total,
        "active":         active,
        "inactive":       total - active,
        "expiring_soon":  expiring_soon,
        "po_at_risk":     po_at_risk,
        "clients_served": clients_served,
    }


def create(db: Session, payload: ConsultantCreate) -> HRBPConsultant:
    if db.query(HRBPConsultant).filter_by(emp_id=payload.emp_id).first():
        raise HTTPException(
            status_code=409,
            detail=f"Employee ID '{payload.emp_id}' already exists",
        )
    record = HRBPConsultant(**payload.model_dump())
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


def list_paginated(
    db: Session,
    page_no: int,
    per_page: int,
    hrbp_ids: list[int] | None = None,
    bh_client_ids: list[int] | None = None,
    client_id: int | None = None,
    cohort: str | None = None,
    perf_tier: str | None = None,
    is_active: bool | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> PageResult:
    q = db.query(HRBPConsultant)
    # hrbp sees their own consultants; bh sees consultants via their client_ids
    if hrbp_ids is not None:
        q = q.filter(HRBPConsultant.hrbp_id.in_(hrbp_ids))
    elif bh_client_ids is not None:
        q = q.filter(HRBPConsultant.client_id.in_(bh_client_ids))
    if client_id is not None:
        q = q.filter(HRBPConsultant.client_id == client_id)
    if cohort is not None:
        q = q.filter(HRBPConsultant.cohort == cohort)
    if perf_tier is not None:
        q = q.filter(HRBPConsultant.perf_tier == perf_tier)
    if is_active is not None:
        q = q.filter(HRBPConsultant.is_active == is_active)
    if date_from is not None:
        q = q.filter(HRBPConsultant.created_at >= date_from)
    if date_to is not None:
        q = q.filter(HRBPConsultant.created_at <= date_to)
    q = q.order_by(HRBPConsultant.name)
    return paginate(q, page_no, per_page)


def get_by_id(db: Session, id: int) -> HRBPConsultant:
    record = db.query(HRBPConsultant).filter_by(id=id).first()
    if not record:
        raise HTTPException(status_code=404, detail="Consultant not found")
    return record


def get_by_emp_id(db: Session, emp_id: str) -> HRBPConsultant:
    record = db.query(HRBPConsultant).filter_by(emp_id=emp_id).first()
    if not record:
        raise HTTPException(status_code=404, detail=f"Employee ID '{emp_id}' not found")
    return record


def update(db: Session, id: int, payload: ConsultantUpdate) -> HRBPConsultant:
    record = get_by_id(db, id)
    update_data = payload.model_dump(exclude_unset=True)

    # Skip po_risk update if the consultant has any open/escalated tickets
    if "po_risk" in update_data:
        has_active_ticket = db.execute(
            select(hrbp_ticket_consultants.c.ticket_id)
            .join(HRBPTicket, HRBPTicket.id == hrbp_ticket_consultants.c.ticket_id)
            .where(
                hrbp_ticket_consultants.c.consultant_id == id,
                HRBPTicket.status.in_(["open", "escalated"]),
            )
            .limit(1)
        ).first()
        if has_active_ticket:
            del update_data["po_risk"]

    for field, value in update_data.items():
        setattr(record, field, value)
    db.commit()
    db.refresh(record)
    return record


def delete(db: Session, id: int) -> None:
    record = get_by_id(db, id)
    db.delete(record)
    db.commit()


def _parse_number(val) -> Decimal | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return Decimal(str(val))
    cleaned = str(val).replace(",", "").strip()
    try:
        return Decimal(cleaned)
    except Exception:
        return None


def bulk_upsert_from_excel(db: Session, file_bytes: bytes) -> dict:
    import io
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active

    raw_headers = [str(c.value).strip().lower() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(raw_headers)}

    inserted = 0
    updated = 0
    errors: list[dict] = []
    now = datetime.now(timezone.utc)

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        try:
            raw_emp = row[col["emp_id"]] if "emp_id" in col else None
            emp_id = str(raw_emp).strip() if raw_emp is not None else None
            if not emp_id or emp_id.lower() == "none":
                continue

            def _get(header: str):
                return row[col[header]] if header in col else None

            name_val = _get("name")
            name = str(name_val).strip() if name_val else None

            join_dt = _get("join_date")
            join_date = join_dt.date() if isinstance(join_dt, datetime) else join_dt

            phone_val = _get("phone")
            phone = str(int(phone_val)) if isinstance(phone_val, float) else (str(phone_val).strip() if phone_val else None)

            email_val = _get("email")
            email = str(email_val).strip() if email_val else None

            monthly_po = _parse_number(_get("monthly_po"))

            # ctc in Excel is annual CTC; store as yearly_ctc and derive monthly_ctc
            yearly_ctc = _parse_number(_get("ctc")) or _parse_number(_get("yearly_ctc"))
            monthly_ctc = (yearly_ctc / 12).quantize(Decimal("0.01")) if yearly_ctc else None

            is_active_val = _get("is_active")
            is_active = bool(is_active_val) if is_active_val is not None else True

            skill_val = _get("skill")
            skill = str(skill_val).strip() if skill_val else None

            designation_val = _get("designation")
            designation = str(designation_val).strip() if designation_val else None
            modality = designation  # keep modality in sync with designation

            margin = _parse_number(_get("margin"))

            po_end_dt = _get("po_end_date")
            po_end_date = po_end_dt.date() if isinstance(po_end_dt, datetime) else po_end_dt

            # Excel has typo "cleint_id"
            client_id_val = _get("cleint_id") or _get("client_id")
            client_id = int(client_id_val) if client_id_val is not None else None

            hrbp_raw = _get("hrbp_id")
            hrbp_id = int(hrbp_raw) if hrbp_raw is not None else None

            existing = db.query(HRBPConsultant).filter_by(emp_id=emp_id).first()

            if existing:
                if name:
                    existing.name = name
                if phone:
                    existing.phone = phone
                if email:
                    existing.email = email
                if monthly_po is not None:
                    existing.monthly_po = monthly_po
                if monthly_ctc is not None:
                    existing.monthly_ctc = monthly_ctc
                if yearly_ctc is not None:
                    existing.yearly_ctc = yearly_ctc
                if margin is not None:
                    existing.margin = margin
                if join_date:
                    existing.join_date = join_date
                if po_end_date:
                    existing.po_end_date = po_end_date
                if skill:
                    existing.skill = skill
                if designation:
                    existing.designation = designation
                if modality:
                    existing.modality = modality
                if client_id is not None:
                    existing.client_id = client_id
                if hrbp_id is not None:
                    existing.hrbp_id = hrbp_id
                existing.is_active = is_active
                existing.updated_at = now
                updated += 1
            else:
                # Pre-populate created_at from join_date so historical records look correct
                created_at = (
                    datetime(join_date.year, join_date.month, join_date.day, tzinfo=timezone.utc)
                    if join_date
                    else now
                )
                record = HRBPConsultant(
                    emp_id=emp_id,
                    name=name or emp_id,
                    phone=phone,
                    email=email,
                    monthly_po=monthly_po,
                    monthly_ctc=monthly_ctc,
                    yearly_ctc=yearly_ctc,
                    margin=margin,
                    join_date=join_date,
                    po_end_date=po_end_date,
                    skill=skill,
                    designation=designation,
                    modality=modality,
                    client_id=client_id,
                    hrbp_id=hrbp_id,
                    is_active=is_active,
                    created_at=created_at,
                    updated_at=now,
                )
                db.add(record)
                inserted += 1

        except Exception as exc:
            errors.append({"row": row_num, "error": str(exc)})

    db.commit()
    return {"inserted": inserted, "updated": updated, "errors": errors}
