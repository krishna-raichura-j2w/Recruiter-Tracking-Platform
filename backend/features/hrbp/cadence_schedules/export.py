from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from features.hrbp.cadence_schedules.service import list_all_sessions
from features.hrbp.storage.service import upload_bytes

HEADERS = ["DATE", "CLIENT", "PROJECT", "CONSULTANT", "RAG", "COMMENT / NOTES"]
COL_WIDTHS = [15, 25, 25, 25, 12, 45]

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F6B3A")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=10)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

RAG_COLORS = {
    "red": "FF4C4C",
    "green": "5CB85C",
    "amber": "F0A500",
    "orange": "F0A500",
}


def _rag_fill(value: str | None) -> PatternFill | None:
    if not value:
        return None
    color = RAG_COLORS.get(str(value).lower())
    if color:
        return PatternFill(fill_type="solid", fgColor=color)
    return None


def build_and_upload(
    db: Session,
    hrbp_ids: list[int] | None = None,
    client_id: int | None = None,
    consultant_id: int | None = None,
    status: str | None = None,
    scheduled_date: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    client_ids: list[int] | None = None,
) -> str:
    result = list_all_sessions(
        db,
        page_no=1,
        per_page=-1,
        hrbp_ids=hrbp_ids,
        client_id=client_id,
        consultant_id=consultant_id,
        status=status,
        scheduled_date=scheduled_date,
        date_from=date_from,
        date_to=date_to,
        client_ids=client_ids,
    )
    rows = result["items"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Cadence Registry"

    # ---------- header row ----------
    ws.row_dimensions[1].height = 22
    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ---------- data rows ----------
    for row_idx, item in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 18

        values = [
            str(item.get("scheduled_date", "") or ""),
            item.get("client_name") or "",
            item.get("project_name") or "—",
            item.get("consultant_name") or "",
            item.get("rca_status") or "Pending",
            item.get("comments") or "N/A",
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = ALIGN_CENTER if col_idx != 6 else ALIGN_LEFT

            # colour the RAG cell
            if col_idx == 5:
                fill = _rag_fill(item.get("rca_status"))
                if fill:
                    cell.fill = fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    # ---------- freeze header ----------
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return upload_bytes(buf.read(), "cadence_registry.xlsx")
