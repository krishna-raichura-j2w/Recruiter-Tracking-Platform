#!/usr/bin/env python3
"""
Recruiter sourcing report — for every candidate a recruiter sourced,
shows: Recruiter Name, Job ID (OL), Client Name, Role Title,
        Candidate Name, Candidate Email, Submission Date, Candidate Status.

Usage:
    python scripts/export_recruiter_sourcing.py
    python scripts/export_recruiter_sourcing.py --recruiter "Priya"
    python scripts/export_recruiter_sourcing.py --out /tmp/sourcing.xlsx
"""
import os
import sys
import argparse
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote

# ── load .env ─────────────────────────────────────────────────────────────────
env_file = Path(__file__).parent.parent / ".env"
if env_file.exists():
    for line in env_file.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, val = line.partition("=")
        os.environ.setdefault(key.strip(), val.strip().strip("'\""))

import psycopg2
import psycopg2.extras

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter


# ── DB ────────────────────────────────────────────────────────────────────────

def _parse_url(url: str) -> dict:
    userinfo, _, hostinfo = url[url.index("://") + 3:].rpartition("@")
    host_port, _, dbname  = hostinfo.partition("/")
    host, _, port         = host_port.partition(":")
    user, _, password     = userinfo.partition(":")
    return dict(
        host=host, port=int(port) if port else 5432,
        dbname=dbname.split("?")[0],
        user=unquote(user), password=unquote(password),
        connect_timeout=20,
        sslmode="require" if "rds.amazonaws" in host or "supabase" in host else "prefer",
    )


def get_conn():
    url = os.environ.get("DATABASE_URL")
    if not url:
        sys.exit("ERROR: DATABASE_URL not set.")
    return psycopg2.connect(**_parse_url(url), cursor_factory=psycopg2.extras.RealDictCursor)


# ── query ─────────────────────────────────────────────────────────────────────

SQL = """
SELECT
    r.name                                      AS recruiter_name,
    r.email                                     AS recruiter_email,
    j.job_id                                    AS ol_job_id,
    j.id                                        AS internal_job_id,
    j.client_name,
    j.role_title,
    c.full_name                                 AS candidate_name,
    c.email                                     AS candidate_email,
    c.status                                    AS candidate_status,
    c.sourced_at                                AS sourced_date,
    s.submitted_at                              AS submission_date
FROM candidates c
JOIN users  r  ON r.id  = c.sourced_by_id
JOIN jobs   j  ON j.id  = c.job_id
LEFT JOIN submissions s ON s.candidate_id = c.id
WHERE c.sourced_by_id IS NOT NULL
  AND c.sourced_at   <= NOW()
  {recruiter_filter}
ORDER BY r.name, c.sourced_at DESC
"""

STATUS_LABELS = {
    "sourced":               "Sourced",
    "handed_to_recruiter":   "Handed to Recruiter",
    "call_in_progress":      "Call In Progress",
    "ready_for_validation":  "Ready for Validation",
    "validated":             "Validated",
    "needs_rework":          "Needs Rework",
    "on_hold":               "On Hold",
    "rejected":              "Rejected",
    "submitted_to_client":   "Submitted to Client",
    "interview_stage":       "Interview Stage",
    "offer_rolled_out":      "Offer Rolled Out",
    "joined":                "Joined",
    "backed_out":            "Backed Out",
}

STATUS_COLORS = {
    "sourced":               "EFF6FF",
    "submitted_to_client":   "D1FAE5",
    "interview_stage":       "FEF9C3",
    "offer_rolled_out":      "DCFCE7",
    "joined":                "BBF7D0",
    "rejected":              "FEE2E2",
    "backed_out":            "FEE2E2",
    "validated":             "E0F2FE",
    "needs_rework":          "FEF3C7",
}


def fetch(recruiter_filter: str = "") -> list[dict]:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(SQL.format(recruiter_filter=recruiter_filter))
        return cur.fetchall()
    finally:
        conn.close()


# ── Excel ─────────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
ALT_FILL    = PatternFill("solid", fgColor="F8FAFC")
WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN        = Side(style="thin", color="CBD5E1")
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN  = Alignment(horizontal="left",   vertical="center", wrap_text=True)

HEADERS = [
    "Recruiter Name",
    "Recruiter Email",
    "OL Job ID",
    "Client Name",
    "Role Title",
    "Candidate Name",
    "Candidate Email",
    "Status",
    "Sourced Date",
    "Submission Date",
]

COL_WIDTHS = [22, 30, 12, 24, 28, 24, 32, 22, 16, 16]


def _fmt(dt) -> str:
    if dt is None:
        return "—"
    if hasattr(dt, "strftime"):
        return dt.strftime("%d %b %Y")
    return str(dt)


def build_excel(rows: list[dict], out_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Recruiter Sourcing"

    # ── title ─────────────────────────────────────────────────────────────────
    ws.merge_cells(f"A1:{get_column_letter(len(HEADERS))}1")
    tc = ws["A1"]
    tc.value = f"Recruiter Sourcing Report  ·  {datetime.now().strftime('%d %b %Y, %H:%M')}"
    tc.font  = Font(bold=True, size=13, color="1E3A8A")
    tc.alignment = CENTER
    ws.row_dimensions[1].height = 30

    # ── header row ────────────────────────────────────────────────────────────
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = CENTER; cell.border = BORDER
    ws.row_dimensions[2].height = 22
    ws.freeze_panes = "A3"

    # ── data rows ─────────────────────────────────────────────────────────────
    for ri, row in enumerate(rows, 3):
        is_alt  = (ri % 2 == 0)
        base_fill = ALT_FILL if is_alt else WHITE_FILL

        status_key   = (row.get("candidate_status") or "").lower()
        status_label = STATUS_LABELS.get(status_key, status_key.replace("_", " ").title())
        status_color = STATUS_COLORS.get(status_key, "F8FAFC")

        ol_job_id = row.get("ol_job_id") or row.get("internal_job_id") or "—"
        recruiter  = row.get("recruiter_name") or "—"

        row_data = [
            recruiter,
            row.get("recruiter_email") or "—",
            ol_job_id,
            row.get("client_name") or "—",
            row.get("role_title") or "—",
            row.get("candidate_name") or "—",
            row.get("candidate_email") or "—",
            status_label,
            _fmt(row.get("sourced_date")),
            _fmt(row.get("submission_date")),
        ]

        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=col, value=val)
            cell.border    = BORDER
            cell.alignment = LEFT_ALIGN

            if col == 1:                            # Recruiter name — bold
                cell.font = Font(bold=True, size=10)
                cell.fill = PatternFill("solid", fgColor="EFF6FF")
            elif col == 8:                          # Status badge
                cell.fill = PatternFill("solid", fgColor=status_color)
                cell.font = Font(bold=True, size=10)
                cell.alignment = CENTER
            elif col == 10 and row.get("submission_date"):  # Submission date highlight
                cell.fill = PatternFill("solid", fgColor="D1FAE5")
                cell.font = Font(size=10)
            else:
                cell.fill = base_fill
                cell.font = Font(size=10)

        ws.row_dimensions[ri].height = 17

    # ── column widths ─────────────────────────────────────────────────────────
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Summary sheet: per-recruiter counts ───────────────────────────────────
    from collections import defaultdict, Counter
    ws2 = wb.create_sheet("By Recruiter")

    # aggregate
    rec_stats: dict[str, dict] = defaultdict(lambda: {
        "total": 0, "submitted": 0, "joined": 0, "rejected": 0,
        "interview": 0, "offer": 0,
    })
    for row in rows:
        name = row.get("recruiter_name") or "Unknown"
        rec_stats[name]["total"] += 1
        s = (row.get("candidate_status") or "").lower()
        if s == "submitted_to_client":   rec_stats[name]["submitted"]  += 1
        elif s == "interview_stage":     rec_stats[name]["interview"]  += 1
        elif s == "offer_rolled_out":    rec_stats[name]["offer"]      += 1
        elif s == "joined":              rec_stats[name]["joined"]     += 1
        elif s in ("rejected", "backed_out"): rec_stats[name]["rejected"] += 1

    sum_headers = ["Recruiter", "Total Sourced", "Submitted", "Interview", "Offer", "Joined", "Rejected/Backed Out"]
    ws2.merge_cells(f"A1:{get_column_letter(len(sum_headers))}1")
    ws2["A1"].value     = "Sourcing Summary by Recruiter"
    ws2["A1"].font      = Font(bold=True, size=12, color="1E3A8A")
    ws2["A1"].alignment = CENTER
    ws2.row_dimensions[1].height = 26

    for col, h in enumerate(sum_headers, 1):
        cell = ws2.cell(row=2, column=col, value=h)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = CENTER; cell.border = BORDER
    ws2.row_dimensions[2].height = 22
    ws2.freeze_panes = "A3"

    for ri, (name, stats) in enumerate(sorted(rec_stats.items()), 3):
        row_fill = ALT_FILL if ri % 2 == 0 else WHITE_FILL
        vals = [name, stats["total"], stats["submitted"], stats["interview"],
                stats["offer"], stats["joined"], stats["rejected"]]
        for col, val in enumerate(vals, 1):
            cell = ws2.cell(row=ri, column=col, value=val)
            cell.border = BORDER
            cell.font   = Font(bold=(col == 1), size=10)
            cell.fill   = row_fill
            cell.alignment = CENTER if col > 1 else LEFT_ALIGN

    # totals row
    tr = 3 + len(rec_stats)
    totals = ["TOTAL",
              sum(s["total"] for s in rec_stats.values()),
              sum(s["submitted"] for s in rec_stats.values()),
              sum(s["interview"] for s in rec_stats.values()),
              sum(s["offer"] for s in rec_stats.values()),
              sum(s["joined"] for s in rec_stats.values()),
              sum(s["rejected"] for s in rec_stats.values())]
    for col, val in enumerate(totals, 1):
        cell = ws2.cell(row=tr, column=col, value=val)
        cell.font   = Font(bold=True, size=10)
        cell.fill   = PatternFill("solid", fgColor="1E3A8A")
        cell.font   = Font(bold=True, color="FFFFFF", size=10)
        cell.border = BORDER
        cell.alignment = CENTER if col > 1 else LEFT_ALIGN

    ws2.column_dimensions["A"].width = 24
    for i in range(2, len(sum_headers) + 1):
        ws2.column_dimensions[get_column_letter(i)].width = 16

    # ── save ──────────────────────────────────────────────────────────────────
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"✅  Saved → {out_path}  ({len(rows)} rows, {len(rec_stats)} recruiters)")


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Recruiter sourcing report")
    parser.add_argument("--recruiter", default=None,
                        help="Filter by recruiter name (partial, case-insensitive)")
    parser.add_argument("--out", default=None, help="Output .xlsx path")
    args = parser.parse_args()

    recruiter_filter = ""
    if args.recruiter:
        recruiter_filter = f"AND r.name ILIKE '%{args.recruiter}%'"

    out_path = (
        Path(args.out) if args.out
        else Path(__file__).parent.parent / "exports" /
             f"recruiter_sourcing_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    print("Connecting to database…")
    rows = fetch(recruiter_filter)
    print(f"Fetched {len(rows)} sourcing records.")
    if not rows:
        print("Nothing to export.")
        return
    build_excel(rows, out_path)


if __name__ == "__main__":
    main()
